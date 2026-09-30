from pathlib import Path
import re, sys, math
from collections import defaultdict
import openpyxl
sys.stdout.reconfigure(encoding="utf-8")

src = Path(r"D:\嗜酸氧化亚铁硫杆菌\01_原始资料\2024_Khaleque\from2024-mmc1.xlsx")
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
s1 = wb["Table S1"]
s2 = wb["Table S2"]

def parse_formula(formula):
    text = str(formula or "").strip()
    m = re.search(r"(<=>|<->|=>|->|→|↔|=)", text)
    if not m:
        return {}, f"missing arrow: {text}"
    left, right = text[:m.start()], text[m.end():]
    out = defaultdict(float)
    term_re = re.compile(r"^\s*(?:(\d+(?:\.\d+)?(?:[eE][+-]?\d+)?)\s+)?(.+?)\s*$")
    for side, sign in ((left, -1.0), (right, 1.0)):
        for term in side.split("+"):
            term = term.strip()
            if not term:
                continue
            tm = term_re.match(term)
            if not tm:
                return {}, f"bad term: {term}"
            coeff = float(tm.group(1) or 1.0)
            met = tm.group(2).strip()
            if not re.match(r"^.+\[[A-Za-z0-9_]+\]$", met):
                return {}, f"bad met: {met}"
            out[met] += sign * coeff
    return {k:v for k,v in out.items() if abs(v)>1e-15}, ""

reactions=[]
stoich={}
parse_warnings=[]
for row in s1.iter_rows(min_row=2, values_only=True):
    rid=row[0]
    if not rid: continue
    st,w=parse_formula(row[2])
    reactions.append(rid); stoich[rid]=st
    if w: parse_warnings.append((rid,w,row[2]))
wb.close()

# Use the source workbook again for S3/S4 values.
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
row_data={}
for sheet_name in ("Table S3","Table S4"):
    ws=wb[sheet_name]
    rows={}
    for excel_row in range(3,634):
        vals=list(next(ws.iter_rows(min_row=excel_row,max_row=excel_row,values_only=True)))
        rows[excel_row]=vals
    row_data[sheet_name]=rows
wb.close()

original_tail=["ZZ_TreT2_UDP","ZZ_TreT1_ADP","ZZ_TreYZ","ZZ_Tre_amy","ZZ_Tre_cga","ZZ_Tre_EX","ZZ_glc_transport","ZZ_glc_EX","ZZ_glycogen_EX","ZZ_g1p_transport","ZZ_g1p_EX","ZZ_Tre_transport"]
swap_tail=["ZZ_g1p_transport","ZZ_g1p_EX","ZZ_Tre_transport","ZZ_Tre_amy","ZZ_Tre_cga","ZZ_Tre_EX","ZZ_glc_transport","ZZ_glc_EX","ZZ_glycogen_EX","ZZ_TreT1_ADP","ZZ_TreT2_UDP","ZZ_TreYZ"]
assert set(original_tail)==set(swap_tail)
base_ids=[]
wb = openpyxl.load_workbook(src, data_only=True, read_only=True)
for row in wb["Table S3"].iter_rows(min_row=3, max_row=633, values_only=True):
    base_ids.append(row[0])
wb.close()
assert len(base_ids)==631
assert base_ids[-12:]==original_tail
original_ids=base_ids
swap_ids=base_ids[:-12]+swap_tail
assert len(reactions)==631

flux_cols=[(4,5,6),(9,10,11),(14,15,16),(19,20,21)]
scenario_names=["TreYZ_最大生长","TreYZ_最大产物","TreYZ_30%生长","TreT-ADP_最大生长","TreT-ADP_最大产物","TreT-ADP_30%生长","TreT-UDP_最大生长","TreT-UDP_最大产物","TreT-UDP_30%生长","Ectoine_最大生长","Ectoine_最大产物","Ectoine_30%生长"]

def calc(sheet_name, mapping, internal_only=False):
    rows=row_data[sheet_name]
    flux_by_scenario=[]
    for group in flux_cols:
        for j in range(3):
            f={rid:0.0 for rid in reactions}
            for excel_row, rid in zip(range(3,634), mapping):
                val=rows[excel_row][group[j]]
                f[rid]=float(val or 0.0)
            flux_by_scenario.append(f)
    mets=set()
    for st in stoich.values(): mets.update(st)
    results=[]
    for scen, f in zip(scenario_names, flux_by_scenario):
        residual={m:sum(stoich[r].get(m,0.0)*v for r,v in f.items()) for m in mets}
        if internal_only:
            residual={m:v for m,v in residual.items() if not m.endswith("[e]")}
        ranked=sorted(((abs(v),m,v) for m,v in residual.items()), reverse=True)
        results.append((scen, ranked[0][0], ranked[:5]))
    return results

for sheet in ("Table S3","Table S4"):
    for label,mapping in (("ORIGINAL",original_ids), ("SWAP_FIRST_LAST3",swap_ids)):
        print(f"===== {sheet} {label} all_metabolites =====")
        for scen,mx,top in calc(sheet,mapping,False):
            print(scen, f"{mx:.16g}", top[:3])
        print(f"===== {sheet} {label} no_external_e =====")
        for scen,mx,top in calc(sheet,mapping,True):
            print(scen, f"{mx:.16g}", top[:3])
print("parse_warnings", len(parse_warnings), parse_warnings[:5])

import itertools
import numpy as np

tail_rows = list(range(622, 631))
fixed_path_map = {631: "ZZ_TreT1_ADP", 632: "ZZ_TreT2_UDP", 633: "ZZ_TreYZ"}
remaining_ids = [
    "ZZ_Tre_amy", "ZZ_Tre_cga", "ZZ_Tre_EX", "ZZ_glc_transport", "ZZ_glc_EX",
    "ZZ_glycogen_EX", "ZZ_g1p_transport", "ZZ_g1p_EX", "ZZ_Tre_transport",
]
all_mets = sorted(set().union(*[set(st) for st in stoich.values()]))
met_index = {m:i for i,m in enumerate(all_mets)}

def scenario_flux_vector(sheet_name, excel_row):
    vals = row_data[sheet_name][excel_row]
    out=[]
    for group in flux_cols:
        for j in range(3): out.append(float(vals[group[j]] or 0.0))
    return out

def fixed_residual(sheet_name):
    res = np.zeros((12, len(all_mets)), dtype=float)
    for s in range(12):
        f={rid:0.0 for rid in reactions}
        # Fixed portion of the vector, including the three pathway reactions.
        for excel_row in range(3, 634):
            if excel_row in tail_rows:
                continue
            rid = fixed_path_map.get(excel_row, base_ids[excel_row-3])
            f[rid] = scenario_flux_vector(sheet_name, excel_row)[s]
        for rid, st in stoich.items():
            v=f[rid]
            if v:
                for met, coef in st.items(): res[s, met_index[met]] += coef*v
    return res

def search_sheet(sheet_name):
    base = fixed_residual(sheet_name)
    touched = sorted(set().union(*[set(stoich[rid]) for rid in remaining_ids]))
    ti = np.array([met_index[m] for m in touched], dtype=int)
    base_t = base[:, ti]
    # contribution[row_index, candidate_reaction, scenario, touched_metabolite]
    contrib = np.zeros((len(tail_rows), len(remaining_ids), 12, len(touched)), dtype=float)
    for ir, excel_row in enumerate(tail_rows):
        flux = scenario_flux_vector(sheet_name, excel_row)
        for ic, rid in enumerate(remaining_ids):
            for s, v in enumerate(flux):
                for met, coef in stoich[rid].items():
                    if met in touched:
                        contrib[ir, ic, s, touched.index(met)] = coef*v
    perms = list(itertools.permutations(range(len(remaining_ids))))
    best=(float("inf"), None)
    exact=[]
    for start in range(0, len(perms), 2000):
        p=np.asarray(perms[start:start+2000], dtype=int)
        c=contrib[np.arange(len(tail_rows))[:,None], p.T, :, :]
        # c shape is (9, batch, 12, nmet), sum over tail rows.
        total=base_t[None,:,:] + np.transpose(c, (1,0,2,3)).sum(axis=1)
        score=np.max(np.abs(total), axis=(1,2))
        k=int(np.argmin(score))
        if float(score[k]) < best[0]:
            best=(float(score[k]), p[k].tolist())
        for j,sc in enumerate(score):
            if float(sc) <= 1e-8: exact.append((float(sc), p[j].tolist()))
    print(f"===== PERMUTATION SEARCH {sheet_name} =====")
    print("best_max_abs_selected_met", best[0])
    print("best_map", {row:remaining_ids[i] for row,i in zip(tail_rows,best[1])})
    print("exact_count", len(exact), "first", exact[:3])
    if best[1] is not None:
        full = base_ids.copy()
        for row,rid in fixed_path_map.items(): full[row-3]=rid
        for row,i in zip(tail_rows,best[1]): full[row-3]=remaining_ids[i]
        for scen,mx,top in calc(sheet_name, full, False):
            print("BEST", scen, f"{mx:.16g}", top[:3])

search_sheet("Table S3")
search_sheet("Table S4")
