import collections
from pathlib import Path
from openpyxl import load_workbook

p=Path('outputs/G53/G53_逐基因证据审查_阶段稿.xlsx')
w=load_workbook(p,read_only=True,data_only=True)
idx=list(w['G53基因索引'].values)[1:]
review=list(w['2026-extended from 2016'].values)[1:]
sources=list(w['2026-61 papers sources'].values)[1:]
audit=list(w['G53分段核查'].values)[1:]
assert len(idx)==293
assert [r[0] for r in idx]==list(range(587,880))
assert len({r[1] for r in idx})==293
assert all(r[9]=='一致' for r in idx)
assert len(audit)==293 and [r[0] for r in audit]==list(range(587,880))
gene=[r for r in review if r[16]=='gene-review']
rxn=[r for r in review if r[16]=='reaction-review']
assert len(gene)==293 and len(rxn)==73
assert collections.Counter(r[22] for r in gene)==collections.Counter(r[1] for r in idx)
source_ids={r[0] for r in sources}
assert len(source_ids)==len(sources)
assert all(set(str(r[22]).split(';'))<=source_ids for r in idx)
assert all(set(str(r[28]).split(';'))<=source_ids for r in review)
original=load_workbook('A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx',read_only=True,data_only=True)
orig={tuple(row[:16]) for row in list(original['Table 1'].values)[2:] if row[0]}
assert all(tuple(r[:16]) in orig for r in rxn)
assert all(r[18] in {'证据不足暂缓','仅更新证据或编号（待确认）','无需修改（本轮未见代谢反应）'} for r in idx)
print({'gene_count':len(idx),'reaction_review_rows':len(rxn),'gff_consistent':sum(r[9]=='一致' for r in idx),'agent_rows':len(audit),'source_rows':len(sources),'model_A_to_P_exact':len(rxn),'actions':dict(collections.Counter(r[18] for r in idx))})
