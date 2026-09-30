import csv
import re
import urllib.request
import urllib.parse
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
source = ROOT / 'G5X_2929基因逐基因证据链' / 'G510_原始清单.txt'
genes = []
pat = re.compile(r'^(\d+)｜(RU820_RS\d+)｜(WP_\d+\.\d+)｜(NZ_CP136162\.1:\d+-\d+\([+-]\))｜(.+)$')
for line in source.read_text(encoding='utf-8').splitlines():
    m=pat.match(line)
    if m: genes.append(m.groups())
print('genes',len(genes),genes[0],genes[-1])

mapping=defaultdict(list)
with (ROOT/'B1_最新基因组与旧AFE编号映射'/'标准化数据'/'旧新基因映射.tsv').open(encoding='utf-8-sig',newline='') as f:
    for row in csv.DictReader(f,delimiter='\t'):
        if row['当前GCF049位点']: mapping[row['当前GCF049位点']].append(row)
print('mapping_status',Counter(tuple(x['映射状态'] for x in mapping[g[1]]) if mapping[g[1]] else ('none',) for g in genes))

w=load_workbook(ROOT/'A0_复现2016与2024模型'/'A0_主任务'/'脚本与环境'/'中间'/'mmc1.xlsx',read_only=True,data_only=True)
rows=[]
for row in w['Table 1'].iter_rows(min_row=3,values_only=True):
    if row[0]: rows.append(row)
byafe=defaultdict(list)
for row in rows:
    blob=' '.join(str(x or '') for x in row[7:10])
    for afe in set(re.findall(r'AFE_\d{4}',blob)):
        byafe[afe].append(row)
linked=[]
for g in genes:
    hits=[]
    for m in mapping[g[1]]:
        hits.extend(byafe.get(m['旧AFE位点'],[]))
    if hits: linked.append((g[0],g[1],g[4],','.join(sorted(set(x[0] for x in hits)))))
print('mapped genes',sum(bool(mapping[g[1]]) for g in genes),'2016-linked genes',len(linked),'reaction rows',sum(len(byafe.get(m['旧AFE位点'],[])) for g in genes for m in mapping[g[1]]))
print('linked sample')
for x in linked: print('\t'.join(x))

cache=ROOT/'G5X_2929基因逐基因证据链'/'G510_UniProt_243159_20260924.tsv'
if not cache.exists():
    url='https://rest.uniprot.org/uniprotkb/stream?'+urllib.parse.urlencode({
        'query':'organism_id:243159','format':'tsv',
        'fields':'accession,reviewed,protein_name,gene_names,ec,xref_refseq,xref_interpro,xref_pfam,xref_kegg'
    })
    cache.write_bytes(urllib.request.urlopen(url,timeout=60).read())
print('uniprot bytes',cache.stat().st_size)
kegg=ROOT/'G5X_2929基因逐基因证据链'/'G510_KEGG_afr_KO_20260924.tsv'
if not kegg.exists():
    kegg.write_bytes(urllib.request.urlopen('https://rest.kegg.jp/link/ko/afr',timeout=30).read())
print('kegg bytes',kegg.stat().st_size)
