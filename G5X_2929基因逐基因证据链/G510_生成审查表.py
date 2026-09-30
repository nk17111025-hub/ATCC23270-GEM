import csv
import re
from urllib.parse import unquote
from collections import Counter, defaultdict
from pathlib import Path
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
HERE = Path(__file__).resolve().parent
OUT = HERE / 'G510_交付'
OUT.mkdir(exist_ok=True)

def read_tsv(path):
    with path.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f, delimiter='\t'))

def write_tsv(path, headers, rows):
    with path.open('w', encoding='utf-8-sig', newline='') as f:
        writer=csv.DictWriter(f, fieldnames=headers, delimiter='\t', extrasaction='ignore')
        writer.writeheader(); writer.writerows(rows)

pat=re.compile(r'^(\d+)｜(RU820_RS\d+)｜(WP_\d+\.\d+)｜(NZ_CP136162\.1):(\d+)-(\d+)\(([+-])\)｜(.+)$')
genes=[]
for line in (HERE/'G510_原始清单.txt').read_text(encoding='utf-8').splitlines():
    m=pat.match(line)
    if m:
        n,locus,wp,chrom,start,end,strand,product=m.groups()
        genes.append(dict(n=int(n),locus=locus,wp=wp,chrom=chrom,start=int(start),end=int(end),strand=strand,product=product))
assert len(genes)==292 and [g['n'] for g in genes]==list(range(2638,2930))
assert len({g['locus'] for g in genes})==292

gff={}
path=ROOT/'B1_最新基因组与旧AFE编号映射'/'标准化数据'/'GCF_049532655.1'/'genomic.gff'
for line in path.open(encoding='utf-8'):
    if line.startswith('#'): continue
    fields=line.rstrip('\n').split('\t')
    if len(fields)<9 or fields[2] not in ('gene','CDS'): continue
    attrs=dict(x.split('=',1) for x in fields[8].split(';') if '=' in x)
    locus=attrs.get('locus_tag','')
    if locus not in {g['locus'] for g in genes}: continue
    gff.setdefault(locus,{})[fields[2]]=(fields,attrs)

mapping=defaultdict(list)
for m in read_tsv(ROOT/'B1_最新基因组与旧AFE编号映射'/'标准化数据'/'旧新基因映射.tsv'):
    if m['当前GCF049位点']: mapping[m['当前GCF049位点']].append(m)

uniprot_by_wp=defaultdict(list)
for u in read_tsv(HERE/'G510_UniProt_243159_20260924.tsv'):
    for wp in re.findall(r'WP_\d+\.\d+',u.get('RefSeq','')):
        uniprot_by_wp[wp].append(u)

ko_by_afe=defaultdict(set)
for line in (HERE/'G510_KEGG_afr_KO_20260924.tsv').read_text(encoding='utf-8').splitlines():
    gene,ko=line.split('\t')
    ko_by_afe[gene.removeprefix('afr:')].add(ko.removeprefix('ko:'))

biocyc_base=ROOT/'C2_外部数据库与注释来源核验'/'C2-1_BioCyc与数据库版本核验'/'标准化数据'
bio_rxn_by_rs=defaultdict(list)
for b in read_tsv(biocyc_base/'reactions.tsv'):
    for rs in set(re.findall(r'AFE_RS\d+',b.get('酶',''))):
        bio_rxn_by_rs[rs].append(b)

w=load_workbook(ROOT/'A0_复现2016与2024模型'/'A0_主任务'/'脚本与环境'/'中间'/'mmc1.xlsx',read_only=True,data_only=True)
rxn_by_afe=defaultdict(list)
for r in w['Table 1'].iter_rows(min_row=3,values_only=True):
    if not r[0]: continue
    for afe in set(re.findall(r'AFE_\d{4}', ' '.join(str(x or '') for x in r[7:10]))):
        rxn_by_afe[afe].append(r)

nonmetabolic=re.compile(r'hypothetical|transposase|integrase|transcription|DNA |RNA |ribosomal|ribosome|chaperone|response regulator|restriction|toxin|antitoxin|partition|comEA|Holliday|sigma-|resolvase|maturase|repair|secretory system|protein translocase|LPS export|lipopolysaccharide transport|LPS-assembly|LPS translocon|efflux|Spy/CpxP|CbpM|HtpG|HPr family phosphocarrier',re.I)
transport=re.compile(r'transport|permease|receptor|channel|efflux|TonB|ExbB|ExbD|MFS |PTS sugar|CorA',re.I)
ambiguous=re.compile(r'hypothetical|family|domain-containing|ATP-binding protein|oxidoreductase|hydrolase|transferase|MFS transporter|TonB-dependent receptor',re.I)

alpha_headers=['Reaction ID','Reaction Name','Reaction Formula','Confidence Level','EC Number','PMID','Subsystem','Gene-Reaction Association','Gene-Protein-Reaction Association','Protein-Reaction-Association','Fe²⁺ lb','Fe²⁺ ub','tetrathionate lb','tetrathionate ub','sulfur lb','sulfur ub','Record class','Candidate ID','Candidate group','2016 linked Reaction ID','2026 reaction object','Legacy AFE locus','Current locus','Primary evidence DOI','Evidence type','Evidence boundary','2026 action','2026 reaction score','Source ID','Score scope']
gene_rows=[]; alpha=[]; review_rows=[]
for g in genes:
    locus=g['locus']; wp=g['wp']; maps=mapping[locus]
    gene_feature,cds=gff[locus]['gene'],gff[locus]['CDS']
    f,a=cds
    official_ok=(f[0]==g['chrom'] and int(f[3])==g['start'] and int(f[4])==g['end'] and f[6]==g['strand'] and a.get('protein_id')==wp and unquote(a.get('product',''))==g['product'])
    if not official_ok: raise AssertionError((g,f,a.get('protein_id'),a.get('product')))
    afes=sorted({m['旧AFE位点'] for m in maps if m['映射状态']=='完全匹配'})
    refs=sorted({m['旧RefSeq位点'] for m in maps if m['映射状态']=='完全匹配'})
    up=uniprot_by_wp[wp]
    kos=sorted({ko for afe in afes for ko in ko_by_afe[afe]})
    bios={b['对象ID']:b for rs in refs for b in bio_rxn_by_rs[rs]}
    old={r[0]:r for afe in afes for r in rxn_by_afe[afe]}
    p=g['product']
    if old:
        action='仅更新证据或编号'
        issue='2016 反应有公式；需区分历史反应分数与本轮基因专属证据。'
    elif nonmetabolic.search(p) and not transport.search(p):
        action='无需修改'
        issue='未见可纳入代谢模型的精确小分子反应；该功能属调控、维护、装配或未知。'
    else:
        action='证据不足暂缓'
        issue='未在 2016 GPR 找到本基因；底物/产物、区室、方向或复合体尚未逐项证实。'
    if not afes:
        issue+=' 旧 AFE 未确认，不推定历史 GPR。'
    if 2647<=g['n']<=2648:
        action='证据不足暂缓'; issue='FE32tpp 的旧 GPR 含该位点，但 TonB receptor/MarC 注释不能证实 Fe³⁺ 专属转运或全部 AND 成员；先做独立 GPR QC。'
    if g['n']==2789:
        action='证据不足暂缓'; issue='CYT2 旧 Gene-Reaction Association 与 Gene-Protein-Reaction Association 不一致；旧式先保留，待复核原表与序列。'
    if 2906<=g['n']<=2909:
        issue='HYD1pp 旧式有 HynS/Isp1/Isp2/HynL；ATCC 23270 整细胞 H₂ 利用已测，quinone 耦联、膜侧和单亚基催化仍未由同株直接测得。'
    if 2846<=g['n']<=2849:
        issue='四个 CorA 同源位点并存；需分开验证 Mg²⁺/其他二价离子选择性、方向和实际表达。'
    if g['n'] in (2638,2639,2640,2641,2642,2643):
        issue='邻近 ABC 组件提示肽摄取系统；未确认肽种类、ATP 系数、跨膜侧与必需亚基组合。'
    if g['n'] in (2698,2699):
        issue='RUBISCO 已表示多个同工酶；本位点是旧 GPR 的一对大小亚基，方向与反应分数不能据此确认。'
    if g['n'] in range(2830,2838):
        issue='ATPS5rpp 含 5 H⁺/ATP 的旧模型假设；基因仅支持 ATP synthase 亚基，不能证明质子系数。'
    if g['n'] in (2784,2785,2786):
        issue='CYTAA31 旧模型有复合体反应；当前 hypothetical 位点身份与质子/膜侧计量待独立验证。'
    if ambiguous.search(p) and not old:
        issue+=' 自动家族名不足以锁定精确反应。'
    source_ids=['G510-S01','G510-S02','G510-S03','G510-S04','G510-S05']
    if bios: source_ids.append('G510-S06')
    if old: source_ids.append('G510-S07')
    if 2904<=g['n']<=2913: source_ids+=['G510-S08','G510-S09','G510-S10']
    rid=f"G510-G-{g['n']}"
    row=dict(总序号=g['n'],当前locus=locus,当前WP=wp,坐标=f"{g['chrom']}:{g['start']}-{g['end']}({g['strand']})",当前product=p,
             官方GFF_CDS核对='一致' if official_ok else '冲突',旧AFE=';'.join(afes) or '未确认',旧AFE映射依据=';'.join(sorted({m['映射方法'] for m in maps})) or '无可靠映射',
             UniProtKB=';'.join(u['Entry'] for u in up),UniProt审校=';'.join(u['Reviewed'] for u in up),UniProt名称=';'.join(u['Protein names'] for u in up),
             UniProt_EC=';'.join(sorted({e.strip() for u in up for e in u['EC number'].split(';') if e.strip()})),
             InterPro=';'.join(sorted({x for u in up for x in u['InterPro'].split(';') if x})),Pfam=';'.join(sorted({x for u in up for x in u['Pfam'].split(';') if x})),
             KEGG_KO=';'.join(kos),BioCyc旧版RXN=';'.join(sorted(bios)),BioCyc证据代码=';'.join(sorted({b['证据代码'] for b in bios.values()})),
             旧模型反应=';'.join(sorted(old)),是否有具体反应='2016模型有公式' if old else '未确认',动作=action,关联行ID=rid+';'+ ';'.join(f"G510-R-{g['n']}-{rxn}" for rxn in sorted(old)),未决点=issue,
             来源ID=';'.join(source_ids),NCBI链接=f'https://www.ncbi.nlm.nih.gov/protein/{wp}',UniProt链接=';'.join('https://www.uniprot.org/uniprotkb/'+u['Entry'] for u in up),
             KEGG链接=';'.join('https://www.kegg.jp/entry/afr:'+afe for afe in afes),BioCyc链接='https://biocyc.org/organism-summary?object=GCF_000021485' if bios else '',
             原始论文DOI='10.1128/aem.56.9.2922-2923.1990;10.1128/jb.184.8.2081-2087.2002' if 2904<=g['n']<=2913 else '',
             论文边界='同株整细胞生长/氢化酶诱导，不定位至该单基因或精确 q8 反应' if 2904<=g['n']<=2913 else '未找到本基因与精确反应的同株直接实验')
    gene_rows.append(row)
    base={h:'' for h in alpha_headers}
    base.update({'Record class':'gene review','Candidate ID':rid,'Candidate group':'G510','2016 linked Reaction ID':';'.join(sorted(old)),
                 '2026 reaction object':'已见2016式；本轮不新建' if old else '未定义', 'Legacy AFE locus':';'.join(afes) or '未确认','Current locus':locus,
                 'Evidence type':'官方CDS；旧新映射；UniProt/InterPro/Pfam/KEGG自动注释；BioCyc旧版；2016原补表',
                 'Evidence boundary':row['论文边界']+'；自动注释和模型行不构成湿实验。','2026 action':action,'2026 reaction score':'不评分',
                 'Source ID':';'.join(source_ids),'Score scope':'gene review 无反应评分'})
    alpha.append(base)
    for rxn,r in sorted(old.items()):
        r_id=f"G510-R-{g['n']}-{rxn}"
        d=dict(zip(alpha_headers[:16],r[:16]))
        oldscore=r[3]
        d.update({'Record class':'reaction review','Candidate ID':r_id,'Candidate group':'G510','2016 linked Reaction ID':rxn,'2026 reaction object':str(r[2] or ''),
                  'Legacy AFE locus':';'.join(afes),'Current locus':locus,'Evidence type':'2016补表精确反应；同株序列与自动注释；必要时 BioCyc/KEGG 交叉引用',
                  'Evidence boundary':'2016 模型定义了该反应；本轮尚无本基因对精确化学式、区室与方向的直接实验。',
                  '2026 action':action,'2026 reaction score':2,
                  'Source ID':';'.join(source_ids),'Score scope':'旧分数是2016行原值；2026=2 仅为精确模型反应的间接序列/注释支持，不继承给GPR、方向或计量。'})
        alpha.append(d)
        review_rows.append(dict(总序号=g['n'],当前locus=locus,旧AFE=';'.join(afes),反应ID=rxn,反应名=r[1],反应式=r[2],原分数=oldscore,本轮反应分数=2,
                                旧EC=r[4],旧PMID=r[5],旧Subsystem=r[6],旧GeneReaction=r[7],旧GeneProteinReaction=r[8],旧ProteinReaction=r[9],
                                铁_lb=r[10],铁_ub=r[11],四硫代硫酸盐_lb=r[12],四硫代硫酸盐_ub=r[13],硫_lb=r[14],硫_ub=r[15],动作=action,证据边界=d['Evidence boundary'],来源ID=d['Source ID']))

write_tsv(OUT/'G510_逐基因索引.tsv',list(gene_rows[0]),gene_rows)
write_tsv(OUT/'G510_Alpha审查.tsv',alpha_headers,alpha)
write_tsv(OUT/'G510_2016反应逐基因对照.tsv',list(review_rows[0]),review_rows)
print('gene_rows',len(gene_rows),'alpha_rows',len(alpha),'reaction_review_rows',len(review_rows))
print('actions',Counter(r['动作'] for r in gene_rows))
print('uniprot linked',sum(bool(r['UniProtKB']) for r in gene_rows),'kegg',sum(bool(r['KEGG_KO']) for r in gene_rows),'biocyc',sum(bool(r['BioCyc旧版RXN']) for r in gene_rows))
