import json
from pathlib import Path
from openpyxl import load_workbook

DIR=Path(__file__).resolve().parent
data=json.loads((DIR/'g52_precheck.json').read_text(encoding='utf-8'))
byloc={r['当前locus']:r for r in data['records']}
manual=json.loads((DIR/'g52_manual_adjudications.json').read_text(encoding='utf-8'))
sheet=load_workbook('A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx',read_only=True,data_only=True)['Table 1']
old={str(r[0]):list(r[:16]) for r in sheet.iter_rows(min_row=2,values_only=True) if r[0]}

specs=[
 ('RU820_RS01960','PUNP1','GPR候选','核查底物后考虑加入 GPR；暂不确定','G52-S001;G52-S004;G52-S007','adenosine + Pi <=> adenine + alpha-D-ribose-1-phosphate'),
 ('RU820_RS01960','PUNP2','GPR候选','核查底物后考虑加入 GPR；暂不确定','G52-S001;G52-S004;G52-S007','inosine + Pi <=> hypoxanthine + alpha-D-ribose-1-phosphate'),
 ('RU820_RS01960','PUNP3','GPR候选','核查底物后考虑加入 GPR；暂不确定','G52-S001;G52-S004;G52-S007','guanosine + Pi <=> guanine + alpha-D-ribose-1-phosphate'),
 ('RU820_RS01960','PUNP4','GPR候选','核查底物后考虑加入 GPR；暂不确定','G52-S001;G52-S004;G52-S007','xanthosine + Pi <=> xanthine + alpha-D-ribose-1-phosphate'),
 ('RU820_RS02115','MTOD','GPR候选','考虑加入 RU820_RS02115 GPR；先核对方向与电荷','G52-S001;G52-S004;G52-S006;G52-S007','S-methyl-5-thio-D-ribulose-1-phosphate -> 2,3-diketo-5-methylthiopentyl-1-phosphate + H2O'),
 ('RU820_RS02190','PNTK','GPR候选','考虑加入 RU820_RS02190 GPR；仅泛酸底物','G52-S001;G52-S004;G52-S006;G52-S007','ATP + (R)-pantothenate -> ADP + (R)-4-phosphopantothenate + H+'),
 ('RU820_RS02985','UDPGD','GPR候选','考虑加入 RU820_RS02985 OR 分支；不复制反应','G52-S001;G52-S004;G52-S006;G52-S007','UDP-glucose + H2O + 2 NAD+ -> UDP-glucuronate + 2 NADH + 3 H+'),
 ('RU820_RS01770',None,'新反应候选','Hold；先核查模型代谢物和同株底物专一性','G52-S001;G52-S006;G52-S007','S-methyl-5-thioinosine + Pi <=> hypoxanthine + S-methyl-5-thio-D-ribose-1-phosphate'),
 ('RU820_RS01775',None,'新反应候选','Hold；先建立 GSSG 代谢物并配平','G52-S001;G52-S007;G52-S008','GSSG + NADPH + H+ -> 2 GSH + NADP+'),
 ('RU820_RS01780',None,'新反应候选','Hold；ROOH 具体底物未定','G52-S001;G52-S007;G52-S008','2 GSH + ROOH -> GSSG + ROH + H2O'),
 ('RU820_RS01930',None,'新反应候选','Hold；核查 S-lactoylglutathione 和乳酸构型','G52-S001;G52-S003;G52-S006;G52-S007','S-lactoylglutathione + H2O <=> GSH + lactate'),
 ('RU820_RS01955',None,'新反应候选','Hold；核查 methylglyoxal 与 lactoylglutathione 代谢物','G52-S001;G52-S003;G52-S006;G52-S007','GSH + methylglyoxal <=> S-lactoylglutathione'),
 ('RU820_RS02225',None,'新反应候选','Hold；沿用 Alpha B-020，醌种和膜侧未定','G52-S001;G52-S006;G52-S007;G52-S010','NAD(P)H + quinone + H+ -> NAD(P)+ + quinol'),
 ('RU820_RS02285',None,'新反应候选','Hold；核查氨基受体和 PLP 位点','G52-S001;G52-S003;G52-S006;G52-S007','L-alanine + 2-oxoglutarate <=> pyruvate + L-glutamate'),
 ('RU820_RS02575',None,'新反应候选','Hold；KEGG adenosine kinase 与当前家族注释需核对','G52-S001;G52-S006;G52-S007','ATP + adenosine -> ADP + AMP'),
 ('RU820_RS02940',None,'新反应候选','Hold；核查 UDP-glucuronate 底物识别','G52-S001;G52-S006;G52-S007','UDP-glucuronate <=> UDP-D-galacturonate'),
 ('RU820_RS01740',None,'功能冲突候选','Hold；DAG kinase 与 undecaprenol kinase 底物冲突','G52-S001;G52-S003;G52-S006;G52-S007','DAG/undecaprenol + ATP -> corresponding phosphate + ADP'),
 ('RU820_RS02405',None,'功能冲突候选','Hold；gamma carbonic anhydrase 与 UniProt transferase 注释冲突','G52-S001;G52-S006;G52-S007','CO2 + H2O <=> HCO3- + H+'),
 ('RU820_RS02790',None,'功能未定候选','Hold；确认是否具有完整 MprF 催化域','G52-S001;G52-S006;G52-S007','L-lysyl-tRNA + phosphatidylglycerol -> lysylphosphatidylglycerol + tRNA'),
 ('RU820_RS02245',None,'体外反应候选','记录 P16.2 体外 TST 反应；Hold 模型新增，cyanide 生理受体及现行 WP 活性未核','G52-S001;G52-S030','thiosulfate + cyanide -> sulfite + thiocyanate'),
]

rows=[]
for loc,rxn,record_class,action,sources,chemistry in specs:
    gene=byloc[loc]
    source=old.get(rxn) if rxn else None
    core=source[:] if source else ['','','','','','','','','','','','','','','','']
    if rxn and source is None:raise ValueError(rxn)
    candidate_id='G52-'+(rxn or 'NEW')+'-'+loc
    note=manual[loc]
    scope=note['issue']+'；化学边界：'+note['chemistry']
    if loc=='RU820_RS02245':
        scope+='；4 仅限历史 P16.2 体外反应，不转给当前 WP 活性或生理受体。'
    evidence=';'.join(dict.fromkeys(sources.split(';')+note['evidence'].split(';')))
    rows.append({'candidate_id':candidate_id,'locus':loc,'wp':gene['WP'],'afe':gene['旧AFE'],'old_reaction':rxn or '', 'core_2016':core,'class':record_class,'action':action,'sources':evidence,'chemistry':chemistry,'scope':scope})

assert len(rows)==len(specs)
assert len({r['candidate_id'] for r in rows})==len(rows)
(DIR/'g52_candidate_rows_draft.json').write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({'candidate_rows':len(rows),'existing_reaction_gpr':sum(bool(r['old_reaction']) for r in rows),'new_reaction':sum(not r['old_reaction'] for r in rows)},ensure_ascii=False))
