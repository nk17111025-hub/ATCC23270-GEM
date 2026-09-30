import collections
import json
import re
from pathlib import Path
from openpyxl import load_workbook

DIR=Path(__file__).resolve().parent
data=json.loads((DIR/'g52_precheck.json').read_text(encoding='utf-8'))
source=Path('A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx')
book=load_workbook(source,read_only=True,data_only=True)
met={r[0]:(r[2],r[3]) for r in book['Table 2'].iter_rows(min_row=2,values_only=True) if r[0]}
terms=re.compile(r'(?:(\d+(?:\.\d+)?)\s+)?([A-Za-z0-9_.-]+\[[a-z]+\])')
formula_parts=re.compile(r'([A-Z][a-z]?)(\d*)')

def elements(formula):
    if not isinstance(formula,str) or not re.fullmatch(r'(?:[A-Z][a-z]?\d*)+',formula):
        return None
    result=collections.Counter()
    for element,n in formula_parts.findall(formula):result[element]+=int(n or 1)
    return result

out=[]
for rxn in {r['Reaction ID']:r for r in data['reaction_rows']}.values():
    raw=rxn['Reaction Formula']
    arrow=next((x for x in ('<=>','->','<-') if x in raw),None)
    if not arrow:
        out.append({'reaction':rxn['Reaction ID'],'status':'parse error','formula':raw})
        continue
    sides=raw.split(arrow)
    delta=collections.Counter()
    missing=[]
    for sign,side in ((-1,sides[0]),(1,sides[1])):
        for n,abbr in terms.findall(side):
            if abbr not in met:
                missing.append(abbr)
                continue
            formula,charge=met[abbr]
            e=elements(formula)
            if e is None or charge is None:
                missing.append(abbr+' (formula/charge missing)')
                continue
            factor=sign*float(n or 1)
            for el,count in e.items():delta[el]+=factor*count
            delta['charge']+=factor*float(charge)
    diff={k:v for k,v in delta.items() if abs(v)>1e-6}
    status='unknown' if missing else 'unbalanced' if diff else 'balanced'
    out.append({'reaction':rxn['Reaction ID'],'status':status,'difference_products_minus_reactants':diff,'missing':sorted(set(missing)),'formula':raw})

(DIR/'g52_2016_balance.json').write_text(json.dumps(out,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(dict(collections.Counter(r['status'] for r in out)),ensure_ascii=False))
for r in out:
    if r['status']!='balanced':print(r['reaction'],r['status'],r['difference_products_minus_reactants'],r['missing'])
