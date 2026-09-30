import collections
import json
import re
from pathlib import Path
from openpyxl import load_workbook

DIR=Path(__file__).resolve().parent
data=json.loads((DIR/'g52_precheck.json').read_text(encoding='utf-8'))
kegg=json.loads((DIR/'g52_kegg_reaction_entries.json').read_text(encoding='utf-8'))['reactions']
records={r['当前locus']:r for r in data['records']}
book=load_workbook('A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx',read_only=True,data_only=True)
met={r[0]:str(r[6] or '').strip() for r in book['Table 2'].iter_rows(min_row=2,values_only=True) if r[0]}
term=re.compile(r'(?:(\d+(?:\.\d+)?)\s+)?([A-Za-z0-9_.-]+\[[a-z]+\])')
kterm=re.compile(r'(?:(\d+(?:\.\d+)?)\s+)?(C\d{5})')

def model_sides(formula):
    arrow=next((a for a in ('<=>','->','<-') if a in formula),None)
    if not arrow:return None
    result=[]; missing=[]
    for side in formula.split(arrow):
        cs=collections.Counter()
        for n,abbr in term.findall(side):
            kid=met.get(abbr,'')
            if re.fullmatch(r'C\d{5}',kid):cs[kid]+=float(n or 1)
            else:missing.append(abbr)
        result.append(cs)
    return result,sorted(set(missing))

def kegg_sides(equation):
    arrow=next((a for a in ('<=>','=>','<=') if a in equation),None)
    if not arrow:return None
    return [collections.Counter({c:float(n or 1) for n,c in kterm.findall(side)}) for side in equation.split(arrow)]

def score(model,reference):
    a,b=model;c,d=reference
    def sim(x,y):
        keys=set(x)|set(y)
        return sum(min(x[k],y[k]) for k in keys)/sum(max(x[k],y[k]) for k in keys) if keys else 0
    return max((sim(a,c)+sim(b,d))/2,(sim(a,d)+sim(b,c))/2)

results=[]
for row in data['reaction_rows']:
    gene=records[row['Current locus']]
    parsed=model_sides(row['Reaction Formula'])
    if parsed is None:continue
    model,missing=parsed
    matches=[]
    for rid in gene['KEGG live reaction'].split(';'):
        if rid and rid in kegg:
            ref=kegg_sides(kegg[rid].get('EQUATION',''))
            if ref:matches.append({'kegg_reaction':rid,'similarity':round(score(model,ref),4),'definition':kegg[rid].get('DEFINITION','')})
    matches.sort(key=lambda x:x['similarity'],reverse=True)
    results.append({'locus':row['Current locus'],'old_reaction':row['Reaction ID'],'model_compounds':[{k:v for k,v in side.items()} for side in model],'model_missing_kegg_ids':missing,'best':matches[:3]})

(DIR/'g52_2016_kegg_chemistry.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
for x in results:
    if x['best']:
        best=x['best'][0]
        if best['similarity']<0.75 or x['model_missing_kegg_ids']:
            print(x['locus'],x['old_reaction'],'best',best['kegg_reaction'],best['similarity'],'missing',','.join(x['model_missing_kegg_ids']))
print('rows',len(results),'best>=0.95',sum(bool(x['best']) and x['best'][0]['similarity']>=.95 and not x['model_missing_kegg_ids'] for x in results))
