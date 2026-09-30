"""Check the proposed TrxR equation using original 2016 metabolite formulas."""
import json
import re
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

DIR = Path(__file__).resolve().parent
source = Path('A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx')
book = load_workbook(source, read_only=True, data_only=True)
wanted = {'h[c]', 'nadph[c]', 'nadp[c]', 'trdox[c]', 'trdrd[c]'}
met = {r[0]: {'formula': r[2], 'charge': r[3], 'name': r[1]} for r in book['Table 2'].iter_rows(min_row=2, values_only=True) if r[0] in wanted}
assert set(met) == wanted

formula = 'h[c] + nadph[c] + trdox[c] -> nadp[c] + trdrd[c]'
delta = Counter()
for sign, species in ((-1, ['h[c]', 'nadph[c]', 'trdox[c]']), (1, ['nadp[c]', 'trdrd[c]'])):
    for m in species:
        for el, n in re.findall(r'([A-Z][a-z]?)(\d*)', met[m]['formula']):
            delta[el] += sign * int(n or 1)
        delta['charge'] += sign * met[m]['charge']
diff = {k: v for k, v in delta.items() if v}
assert not diff
out = {
    'reaction_id': 'TRDR',
    'legacy_formula': 'nadph[c] + trdrd[c] -> 3 h[c] + nadp[c] + trdox[c]',
    'candidate_formula': formula,
    'metabolites': met,
    'products_minus_reactants': diff,
    'balanced': True,
    'biochemical_support': '10.1007/s00284-009-9390-2',
    'source': '2016 iMC507 Supplementary Table 2 formulas and charges',
    'limitation': 'Full model network and transport/compartment consistency remain to be tested.',
}
(DIR / 'g52_trdr_corrected_balance.json').write_text(json.dumps(out, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(out, ensure_ascii=False))
