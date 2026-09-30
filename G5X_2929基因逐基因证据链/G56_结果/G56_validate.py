import csv
import hashlib
import json
import re
from collections import Counter
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]
data = json.loads((HERE / 'G56_审查数据.json').read_text(encoding='utf-8'))
idx, review, sources = data['index'], data['review'], data['sources']
source_ids = {x[0] for x in sources}

def source_tokens(value):
    return [x for x in str(value or '').split(';') if x]

def normalized(values):
    return [None if x == '' else x for x in values]

assert len(idx) == 293
assert [x['总序号'] for x in idx] == list(range(1466, 1759))
assert len({x['当前locus'] for x in idx}) == 293
assert len({x['WP'] for x in idx}) == 292
assert all(x['GFF身份核对'] == '通过' for x in idx)
assert len(review) == 390
assert Counter(x[16] for x in review) == {
    '2026 gene review': 293,
    '2026 reaction review': 94,
    '2026 reaction candidate—hold': 3,
}
assert len({x[17] for x in review}) == 390
assert {x[22] for x in review if x[16] == '2026 gene review'} == {x['当前locus'] for x in idx}
assert all(x[27] == '不评分' for x in review if x[16] != '2026 reaction review')
assert all(not any(x[:16]) for x in review if x[16] == '2026 gene review')
assert all(set(source_tokens(x['来源ID'])) <= source_ids for x in idx)
assert all(set(source_tokens(x['论文Source ID'])) <= source_ids for x in idx)
assert all(set(source_tokens(x[28])) <= source_ids for x in review)

original = data['model_by_id']
reaction_rows = [x for x in review if x[16] == '2026 reaction review']
assert all(x[:16] == original[x[19]][:16] for x in reaction_rows)
assert all(x[27] == 2 for x in reaction_rows)
assert all(x[26] in {'无需修改', '仅更新证据或编号'} for x in reaction_rows)
changed_scores = [x for x in reaction_rows if x[3] != x[27]]
assert len(changed_scores) == 39
assert len({x[19] for x in reaction_rows}) == 90

book = openpyxl.load_workbook(HERE / 'G56_Alpha格式证据表.xlsx', read_only=True, data_only=True)
assert book.sheetnames == ['2026-extended from 2016', '2026-61 papers sources', 'G56 gene index']
alpha = book.worksheets[0]
assert list(next(alpha.iter_rows(min_row=2, max_row=2, values_only=True))[:30]) == data['alpha_headers']
for row_number, (actual, expected) in enumerate(zip(alpha.iter_rows(min_row=3, values_only=True), review, strict=True), 3):
    assert normalized(actual[:30]) == normalized(expected), (row_number, [(j, actual[j], expected[j]) for j in range(30) if actual[j] != expected[j]])
srcsheet = book.worksheets[1]
assert list(next(srcsheet.iter_rows(values_only=True))) == data['source_headers']
for actual, expected in zip(srcsheet.iter_rows(min_row=2, values_only=True), sources, strict=True):
    assert normalized(actual) == normalized(expected)
ixsheet = book.worksheets[2]
ix_headers = list(idx[0])
assert list(next(ixsheet.iter_rows(values_only=True))) == ix_headers
for actual, expected in zip(ixsheet.iter_rows(min_row=2, values_only=True), idx, strict=True):
    assert normalized(actual) == normalized([expected[k] for k in ix_headers])

gff = ROOT / 'B1_最新基因组与旧AFE编号映射/标准化数据/GCF_049532655.1/genomic.gff'
gff_sha = hashlib.sha256(gff.read_bytes()).hexdigest()
assert gff_sha == '758a8b72423a824840a52e1e4184571474500c7fce155944a28e05fef6577091'

hold = [x for x in idx if x['2026动作'] == '证据不足暂缓']
updates = [x for x in idx if x['2026动作'] == '仅更新证据或编号']
def experiment(x):
    s = x['未决点'] + ' ' + x['当前NCBI product']
    if '映射' in s: return '核实基因组序列、邻域与旧AFE对应关系；必要时核对蛋白序列。'
    if '转运' in s or re.search('transporter|permease|antiporter|symporter|porin', s, re.I):
        return '同株或异源表达背景测底物摄取/外排、方向与能量依赖，并定位膜侧。'
    if 'CsoS3' in s: return '纯化CsoS3测CO2/HCO3−转化，并确定羧酶体定位与模型区室。'
    if 'PPC' in s: return '纯化AFE_1883测PEP羧化；遗传扰动区分AFE_1810的贡献。'
    if 'AHL' in s: return '测定AfeI对各acyl-ACP底物的产物谱与计量，并定位细胞区室。'
    if 'NorD' in s: return '核对NorD/CbbO1蛋白身份；分别测相关酶活或复合体功能。'
    if '酶名' in s or '反应' in s: return '纯化该蛋白或构建定向突变株，测精确底物、产物、计量与方向。'
    return '补充该基因专属的功能实验，明确是否存在可加入模型的代谢反应。'

issue_headers = ['总序号','当前locus','旧AFE','当前product','2026动作','未决点','建议验证实验','论文Source ID','2016关联反应']
with (HERE / 'G56_未决点与验证建议.tsv').open('w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, delimiter='\t')
    w.writerow(issue_headers)
    for x in hold:
        w.writerow([x['总序号'], x['当前locus'], x['旧AFE'], x['当前NCBI product'],
                    x['2026动作'], x['未决点'], experiment(x), x['论文Source ID'], x['2016关联反应']])

with (HERE / 'G56_模型动作清单.tsv').open('w', encoding='utf-8-sig', newline='') as f:
    w = csv.writer(f, delimiter='\t')
    w.writerow(['总序号','当前locus','旧AFE','2026动作','2016关联反应','具体处理','未决点','来源ID'])
    for x in idx:
        treatment = ('只修订本组证据与反应评分；不改模型' if x['2026动作'] == '仅更新证据或编号'
                     else '保留候选，待精确反应/身份核实' if x['2026动作'] == '证据不足暂缓'
                     else '保留现状')
        w.writerow([x['总序号'], x['当前locus'], x['旧AFE'], x['2026动作'],
                    x['2016关联反应'], treatment, x['未决点'], x['来源ID']])

count = Counter(x['2026动作'] for x in idx)
conflict = [x for x in idx if x['数据库分歧'] not in ('', '未见可自动判定的冲突')]
paper_genes = [x for x in idx if x['论文Source ID']]
uniprot_exact = [x for x in idx if x['UniProt accession']]
uniprot_old_only = [x for x in idx if not x['UniProt accession'] and x['UniProt旧AFE关联但WP未匹配']]
report = f'''# G56 全组证据重查与完整性验收

审查范围：ATCC 23270，GCF_049532655.1，G56 总序号 1466–1758；完成于 2026-09-26。逐基因结果见 [Alpha 格式证据表](G56_Alpha格式证据表.xlsx)、[基因索引](G56_基因索引.tsv)、[证据来源](G56_证据来源.tsv)、[模型动作清单](G56_模型动作清单.tsv)、[未决点与验证建议](G56_未决点与验证建议.tsv)。

## 完整性与模型动作

| 检查项 | 结果 |
|---|---:|
| 任务清单基因、独立 locus、独立 WP | 293 / 293 / 292 |
| 官方 GFF 坐标、链向、WP、product 逐项一致 | 293 / 293 |
| 旧 AFE 精确对应 | {sum(x['旧AFE'] != '未确认' for x in idx)} / 293 |
| 旧 AFE 未确认 | {sum(x['旧AFE'] == '未确认' for x in idx)} / 293 |
| 当前 WP 与 UniProt 精确匹配 | {len(uniprot_exact)} / 293 |
| 仅旧 AFE 与 UniProt 相连、当前 WP 未匹配 | {len(uniprot_old_only)} / 293 |
| 有旧模型反应关联的基因 | {sum(bool(x['2016关联反应']) for x in idx)} / 293 |
| 基因审查行 / 2016 已有反应复核行 / 暂缓候选行 | 293 / 94 / 3 |
| 无需修改 / 仅更新证据或编号 / 证据不足暂缓 | {count['无需修改']} / {count['仅更新证据或编号']} / {count['证据不足暂缓']} |
| 本轮正式新增反应、修改 GPR、修改 SBML | 0 / 0 / 0 |

所有 293 个位点均有独立基因审查行和动作。RU820_RS08500 与 RU820_RS08505 共享官方 WP_012606582.1，仍按两个 locus 分开审查。数据库预测、同源及结构域、实验论文分别列在索引末三列；原始论文的直接测量对象和证据边界列在来源表。基因行不评分；只有既有精确反应行给出 2026 评分。旧模型的 A–P 核心列逐格保留。

## 反应级复核

94 行对应 90 个旧反应 ID；4 个反应由本组两个基因分别复核。94 行的 2026 分数均为 2，其中 {len(changed_scores)} 行较 2016 原分数 3 下调，2016 原分数仍保存在 A–P 列。分数只适用于已列明的反应，不能迁移到新基因、新区室或新方向。其余基因及暂缓候选均不评分。

| 位点 / 对象 | 本轮结论 |
|---|---|
| RU820_RS07800 / AFE_1688 CsoS3 | 羧酶体碳酸酐酶候选暂缓；旧 HCO3E 已有关联，需核实区室和该蛋白催化。 |
| RU820_RS07810–07815 / AFE_1690–1691 | RUBISCO 已有反应保留；同株表达和蛋白检出支持功能，未定位该复合体对旧反应的直接酶学测定。 |
| RU820_RS08260 / AFE_1792 | SQRED1 仅更新证据与评分；2025 过表达支持硫氧化生理作用，旧 q8 反应膜侧及 Sre OR 分支未证，三条件上下界全为 0，另作模型质控。 |
| AFE_1906/1907/1908/1910 | 原 PMID 238956 测脂多糖脂肪酸组成，不能逐条证明 Fab 脂肪酸链长反应；相关反应按间接证据评 2。 |
| AFE_1987 / UHGA | 原 PMID 15044494 测 GnnA/GnnB，不能证明 LpxH/UHGA；该反应按注释评 2。 |
| RU820_RS08690 / AFE_1883 | PPC 已有关联 AFE_1810；AFE_1883 同工酶候选继续暂缓，不加入 OR。 |
| RU820_RS09235 / AFE_1999 | AfeI 产 AHL 有同株与异源表达证据；各 acyl-ACP 底物及逐产物计量未定，暂缓精确反应。 |

## 冲突与待验证

索引保留 {len(conflict)} 个带数据库分歧的位点；{len(hold)} 个暂缓位点逐条列在“未决点与验证建议”。需优先处理旧 AFE 映射歧义、NorD/CbbO1 指称、CsoS3 区室、AFE_1883 的 PPC 活性、AFE_1847 转运底物和能量耦联、AfeI 底物谱、SQRED1 的 q8 膜侧及模型全零界限。

本轮识别 {len(paper_genes)} 个带本组论文 Source ID 的基因；论文测得表达、蛋白丰度、表型或产物时，索引均保留直接测量对象与精确反应之间的边界。未取得基因级 TCDB 结果，也未独立取得 MetaCyc 原始记录。UniProt 快照中仍有 {293-len(uniprot_exact)-len(uniprot_old_only)} 个位点既无当前 WP 精确条目，也无仅旧 AFE 关联条目。这些缺口没有转换为肯定反应或高分。

## 可复核来源与自检

使用官方 GFF、B1 映射、D3 数据库交叉引用、本地 iMC507 Alpha 基线、2026-09-26 UniProt taxon 243159 快照、Rhea 反应式及来源表所列原始论文。证据表中的 18 条来源记录分别标明菌株、实测内容和适用边界。

GFF SHA-256：`{gff_sha}`。程序核对通过：293 个清单位点顺序和身份、390 个不重复审查行、18 个有效来源 ID、94 行原始 A–P 列与 Alpha 基线逐格一致、工作簿三个 sheet 与生成数据逐格一致。没有将数据库候选反应号直接写作模型 Reaction ID。
'''
(HERE / 'G56_验收报告.md').write_text(report, encoding='utf-8')
print(json.dumps({
    'genes': len(idx), 'reactions': len(reaction_rows), 'changed_scores': len(changed_scores),
    'holds': len(hold), 'updates': len(updates), 'conflicts': len(conflict),
    'sources': len(sources), 'paper_genes': len(paper_genes),
}, ensure_ascii=False))
