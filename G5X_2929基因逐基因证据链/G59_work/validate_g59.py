"""Read-only identity and structure checks for G59 review workbooks."""

from __future__ import annotations

import argparse
import re
from collections import Counter
from pathlib import Path

import openpyxl


ROOT = Path(__file__).resolve().parents[2]
WORK = Path(__file__).resolve().parent
GFF = ROOT / "B1_最新基因组与旧AFE编号映射/标准化数据/GCF_049532655.1/genomic.gff"
SOURCE = WORK / "G59_source.txt"
ALPHA = ROOT / "00_项目导航与最终成果/ATCC23270-2026_最终模型证据表.xlsx"
MODEL2016 = ROOT / "A0_复现2016与2024模型/A0_主任务/脚本与环境/中间/mmc1.xlsx"


def attributes(text: str) -> dict[str, str]:
    return dict(field.split("=", 1) for field in text.split(";") if "=" in field)


def source_rows() -> list[tuple[int, str, str, str]]:
    rows = []
    pattern = re.compile(
        r"^(\d{4})｜(RU820_RS\d+)｜(WP_\d+\.\d+)｜"
        r"(NZ_CP136162\.1:\d+-\d+\([+-]\))｜"
    )
    for line in SOURCE.read_text(encoding="utf-8-sig").splitlines():
        match = pattern.match(line)
        if match:
            rows.append((int(match[1]), match[2], match[3], match[4]))
    return rows


def official_rows() -> dict[str, tuple[str, str]]:
    genes: dict[str, str] = {}
    cds: dict[str, str] = {}
    with GFF.open(encoding="utf-8") as handle:
        for line in handle:
            if line.startswith("#"):
                continue
            cols = line.rstrip("\n").split("\t")
            if len(cols) != 9:
                continue
            attr = attributes(cols[8])
            locus = attr.get("locus_tag", "")
            if not locus.startswith("RU820_RS"):
                continue
            if cols[2] == "gene" and attr.get("gene_biotype") == "protein_coding":
                genes[locus] = f"{cols[0]}:{cols[3]}-{cols[4]}({cols[6]})"
            elif cols[2] == "CDS":
                cds[locus] = attr.get("protein_id", "")
    return {locus: (wp, coord) for locus, coord in genes.items() if (wp := cds.get(locus))}


def read_parts(paths: list[Path]) -> list[tuple[int, str, str, str, str]]:
    all_rows = []
    alpha_headers = [cell.value for cell in next(openpyxl.load_workbook(ALPHA, read_only=True).worksheets[0].iter_rows(min_row=2, max_row=2))]
    model = openpyxl.load_workbook(MODEL2016, read_only=True, data_only=True)
    model_rows = {row[0]: row for row in model["Table 1"].iter_rows(min_row=3, values_only=True) if row[0]}
    for path in paths:
        book = openpyxl.load_workbook(path, read_only=True, data_only=True)
        index = book.worksheets[0]
        review = next((sheet for sheet in book if "Alpha" in sheet.title), None)
        if review is None:
            raise ValueError(f"{path.name}: Alpha review sheet absent")
        actual_headers = [cell.value for cell in next(review.iter_rows(max_row=1))]
        if len(actual_headers) != 30 or [str(v).strip().lower() for v in actual_headers[:2]] != [str(v).strip().lower() for v in alpha_headers[:2]]:
            raise ValueError(f"{path.name}: Alpha 30-column structure mismatch")
        headers = [str(cell.value or "").lower() for cell in next(index.iter_rows(max_row=1))]
        def find(*terms: str) -> int:
            hits = [i for i, header in enumerate(headers) if all(term in header for term in terms)]
            if len(hits) != 1:
                raise ValueError(f"{path.name}: index header lookup {terms}: {hits}")
            return hits[0]
        number = find("no.") if "total no." in headers else find("总序号")
        locus = find("current", "locus") if any("current locus" in header for header in headers) else find("当前", "locus")
        wp = find("wp")
        coord = find("coord") if any("coord" in header for header in headers) else find("坐标")
        action = headers.index("action") if "action" in headers else headers.index("动作")
        part_rows = []
        for row in index.iter_rows(min_row=2, values_only=True):
            if row[number] is None:
                continue
            part_rows.append((int(row[number]), str(row[locus]), str(row[wp]), str(row[coord]), str(row[action])))
        review_rows = list(review.iter_rows(min_row=2, values_only=True))
        print(path.name, "gene rows", len(part_rows), "review rows", len(review_rows))
        review_loci = Counter(row[22] for row in review_rows if row[22] and "gene" in str(row[16]).lower())
        missing_gene_review = sorted({row[1] for row in part_rows} - review_loci.keys())
        print(path.name, "missing dedicated gene review", missing_gene_review)
        mismatched_model_rows = []
        for row in review_rows:
            if "2016" not in str(row[16]):
                continue
            reaction_id = row[0]
            source_row = model_rows.get(reaction_id)
            if not source_row:
                mismatched_model_rows.append((reaction_id, "missing original ID"))
                continue
            differing = [i + 1 for i in range(16) if row[i] != source_row[i]]
            if differing:
                mismatched_model_rows.append((reaction_id, differing))
        print(path.name, "2016 A:P mismatches", mismatched_model_rows[:20], "total", len(mismatched_model_rows))
        all_rows.extend(part_rows)
        book.close()
    return all_rows


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("parts", nargs="+", type=Path)
    args = parser.parse_args()
    source = source_rows()
    official = official_rows()
    reviewed = read_parts(args.parts)
    assert len(source) == 293, f"source count: {len(source)}"
    assert [row[0] for row in source] == list(range(2345, 2638))
    assert len({row[1] for row in source}) == 293
    expected = {row[0]: row[1:] for row in source}
    counts = Counter(row[0] for row in reviewed)
    bad = []
    for number, locus, wp, coord, action in reviewed:
        if number not in expected or (locus, wp, coord) != expected[number]:
            bad.append((number, locus, wp, coord, "source mismatch"))
        elif (wp, coord) != official.get(locus):
            bad.append((number, locus, wp, coord, "GFF/CDS mismatch"))
        elif not action.strip():
            bad.append((number, locus, wp, coord, "empty action"))
    print("source", len(source), "reviewed", len(reviewed), "unique seq", len(counts))
    print("missing", sorted(set(expected) - counts.keys()))
    print("duplicates", [(number, count) for number, count in counts.items() if count != 1])
    print("mismatches", bad[:20], "total", len(bad))
    assert not bad and len(counts) == 293 and all(count == 1 for count in counts.values())


if __name__ == "__main__":
    main()
