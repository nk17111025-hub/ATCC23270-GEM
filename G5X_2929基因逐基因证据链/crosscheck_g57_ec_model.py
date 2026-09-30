import csv
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
out = root / 'G5X_2929基因逐基因证据链/G57_审查输出'
with (out / 'G57_逐基因索引与状态.tsv').open(encoding='utf-8-sig', newline='') as f:
    genes = list(csv.DictReader(f, delimiter='\t'))

book = load_workbook(root / '00_项目导航与最终成果/ATCC23270-2026_最终模型证据表.xlsx', read_only=True, data_only=True)
model = [r for r in book.worksheets[0].iter_rows(min_row=3, values_only=True) if r[16] == '2016 original']
ec_re = re.compile(r'(?<!\d)(\d+\.\d+\.\d+\.\d+)(?!\d)')
model_by_ec = defaultdict(list)
for r in model:
    for ec in set(ec_re.findall(str(r[4] or ''))):
        model_by_ec[ec].append(r)

rows = []
for g in genes:
    if g['2016 GPR 反应']:
        continue
    ecs = set(ec_re.findall(' '.join((g['UniProt EC'],g['KEGG EC']))))
    for ec in sorted(ecs):
        for r in model_by_ec.get(ec, []):
            rows.append({'总序号':g['总序号'],'RU820':g['RU820'],'WP':g['WP'],'当前product':g['当前 product'],'基因侧EC':ec,'基因侧EC来源':'UniProt;KEGG' if ec in g['UniProt EC'] and ec in g['KEGG EC'] else 'UniProt' if ec in g['UniProt EC'] else 'KEGG','2016同EC反应':r[0],'2016反应名':r[1],'2016反应式':r[2],'2016原GPR':r[7] or '', '证据边界':'EC 分类相同仅为候选；不能据此认定相同底物、辅因子、方向、区室、基因必需性或旧模型遗漏'})

fields = ['总序号','RU820','WP','当前product','基因侧EC','基因侧EC来源','2016同EC反应','2016反应名','2016反应式','2016原GPR','证据边界']
with (out / 'G57_同EC旧模型候选.tsv').open('w', encoding='utf-8-sig', newline='') as f:
    w = csv.DictWriter(f, fieldnames=fields, delimiter='\t')
    w.writeheader()
    w.writerows(rows)
print({'candidate_rows':len(rows),'candidate_genes':len({r['RU820'] for r in rows})})
