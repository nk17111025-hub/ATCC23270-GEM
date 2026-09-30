import csv
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
out = root / 'G5X_2929基因逐基因证据链/G57_审查输出'

def table(name):
    with (out / name).open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f, delimiter='\t'))

genes = table('G57_逐基因索引与状态.tsv')
alpha = table('G57_Alpha逐基因审查.tsv')
reactions = table('G57_Alpha既有反应逐项对照.tsv')
sources = table('G57_来源表.tsv')
triage = table('G57_文献命中逐篇判读.tsv')
primers = table('G57_论文引物身份核验.tsv')
ec_candidates = table('G57_同EC旧模型候选.tsv')
current_search = table('G57_当前编号双路文献检索.tsv')
tcdb_review = table('G57_TCDB家族同源核验.tsv')

assert len(genes) == len(alpha) == 293
assert [int(r['总序号']) for r in genes] == list(range(1759,2052))
assert len({r['RU820'] for r in genes}) == len({r['WP'] for r in genes}) == 293
assert all(r['CDS 坐标核对'] == '通过' for r in genes)
assert len(reactions) == 64 and len({r['Reaction ID'] for r in reactions}) == 33
assert len(sources) == 22 and len(triage) == 17 and len(primers) == 10 and len(current_search) == 293 and len(tcdb_review) == 44
assert sum(r['家族同源支持']=='是' for r in tcdb_review) == 41
assert len(ec_candidates) == 5
source_ids = {r['Source ID'] for r in sources}
for row in genes + alpha + reactions:
    assert set(row['Source ID'].split(';')) <= source_ids
for row in reactions:
    assert row['Reaction ID'] and row['Reaction Formula'] and row['Gene-Reaction Association']

by_id = {int(r['总序号']):r for r in genes}
assert by_id[1759]['RU820'] == 'RU820_RS09280'
assert by_id[2051]['RU820'] == 'RU820_RS10800'
assert 'G57-S19' in by_id[1905]['Source ID']
assert all('G57-S09' in by_id[n]['Source ID'] for n in range(1999,2006))
assert 'G57-S21' in by_id[1804]['Source ID']
assert {r['当前RU820'] for r in primers if r['论文表1引物'].startswith('phn')} == {f'RU820_RS{i}' for i in range(10520,10551,5)}
assert {r['当前RU820'] for r in primers if r['论文表1引物'].startswith('paper_Afe2172')} == {'RU820_RS04185'}

book = load_workbook(root / 'outputs/G57/G57_阶段证据审查.xlsx', read_only=True)
assert len(book.sheetnames) == 12
assert len(list(book['逐基因索引'].rows)) == 294
assert len(list(book['九项证据链'].rows)) == 294
assert len(list(book['Alpha逐基因'].rows)) == 294
assert len(list(book['2016反应对照'].rows)) == 65
assert len(list(book['来源'].rows)) == 23
assert len(list(book['文献逐篇判读'].rows)) == 18
assert len(list(book['论文引物身份'].rows)) == 11
assert len(list(book['同EC旧模型候选'].rows)) == 6
assert len(list(book['当前编号文献检索'].rows)) == 294
assert len(list(book['TCDB家族核验'].rows)) == 45

print({'genes':len(genes),'reactions':len(reactions),'sources':len(sources),'doi_triage':len(triage),'primer_hits':len(primers),'ec_candidates':len(ec_candidates),'current_id_search':len(current_search),'tcdb_queries':len(tcdb_review),'actions':dict(Counter(r['动作'] for r in genes)),'workbook_sheets':len(book.sheetnames),'stage_acceptance':'未通过完整证据链'})
